"""生成技术清单 Word 文档

复用课程设计报告中已验证的排版参数（A4 + 固定表格布局）。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"D:\claude_code\课程\自然语言处理")
OUT = ROOT / "技术清单与论文出处.docx"

CN_FONT = "微软雅黑"
CN_BODY = "宋体"


def set_font(run, name=CN_BODY, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading("", level=level)
    r = h.add_run(text)
    set_font(r, CN_FONT, {1: 15, 2: 13, 3: 11.5}.get(level, 11),
             bold=True, color=RGBColor(0x1A, 0x36, 0x5D))
    return h


def add_field(doc, label: str, value: str, mono: bool = False) -> None:
    """添加一个带标签的字段行"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35

    r = p.add_run(f"{label}：")
    set_font(r, CN_FONT, 10, bold=True, color=RGBColor(0x33, 0x33, 0x33))

    r = p.add_run(value)
    if mono:
        set_font(r, "Consolas", 9.5, color=RGBColor(0x0B, 0x5C, 0x8A))
    else:
        set_font(r, CN_BODY, 10)


def add_link_field(doc, url: str) -> None:
    """添加带超链接的链接行"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35

    r = p.add_run("链接：")
    set_font(r, CN_FONT, 10, bold=True, color=RGBColor(0x33, 0x33, 0x33))

    # 插入真实超链接
    part = p.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "19")
    rPr.append(sz)
    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.text = url
    new_run.append(t)
    hyperlink.append(new_run)
    p._p.append(hyperlink)


def _set_table_layout_fixed(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "auto")
    tblW.set(qn("w:w"), "0")
    tblPr.append(tblW)


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(h)
        set_font(r, CN_FONT, font_size, bold=True)

    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(v))
            set_font(r, CN_BODY, font_size)

    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    _set_table_layout_fixed(t)
    return t


# 数据从 Markdown 生成脚本导入，避免重复维护
import sys
sys.path.insert(0, str(ROOT))
from tools.gen_tech_doc import CATEGORIES  # noqa: E402


def build() -> None:
    doc = Document()

    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)

    # --- 标题 ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("金融领域问答系统 · 技术清单与论文出处")
    set_font(r, CN_FONT, 18, bold=True, color=RGBColor(0x0F, 0x2A, 0x4A))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("《自然语言处理》课程设计    2026 年 9 月")
    set_font(r, CN_BODY, 10, color=RGBColor(0x77, 0x77, 0x77))

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(10)
    intro.paragraph_format.line_spacing = 1.5
    r = intro.add_run(
        "本清单覆盖项目实际使用的全部关键技术，每项标注论文出处、"
        "在本项目中的具体用途与对应代码位置，可直接用于报告参考文献。"
        "引用格式遵循 ACL/ACM 通用规范，链接均经联网核实。"
    )
    set_font(r, CN_BODY, 10)

    # --- 分类展开 ---
    for cat_name, items in CATEGORIES:
        add_heading(doc, cat_name, 2)
        for tech, cite, link, usage, code in items:
            add_heading(doc, tech, 3)
            add_field(doc, "论文出处", cite)
            add_link_field(doc, link)
            add_field(doc, "项目用途", usage)
            add_field(doc, "代码位置", code, mono=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- 速查表 ---
    doc.add_page_break()
    add_heading(doc, "附：技术与项目对应关系速查", 1)
    add_table(
        doc,
        ["层次", "用到的技术", "关键论文"],
        [
            ["生成", "RAG、Transformer、提示工程", "Lewis 2020、Vaswani 2017、Wei 2022"],
            ["嵌入", "SBERT、BGE-M3", "Reimers 2019、Chen 2024"],
            ["检索", "BM25、DPR、HNSW", "Robertson 2009、Karpukhin 2020、Malkov 2020"],
            ["重排", "Cross-Encoder、bge-reranker", "Nogueira 2019"],
            ["融合", "RRF", "Cormack 2009"],
            ["评测", "Recall@k、MRR、FinQA", "Manning 2008、Chen 2021"],
        ],
        widths=[1.8, 6.2, 8.6],
    )

    # --- GB/T 7714 引用格式 ---
    doc.add_paragraph()
    add_heading(doc, "参考文献（GB/T 7714 格式）", 1)

    refs = [
        "[1] CORMACK G V, CLARKE C L A, BUETTCHER S. Reciprocal rank fusion outperforms "
        "condorcet and individual rank learning methods[C]//Proceedings of the 32nd "
        "International ACM SIGIR Conference. New York: ACM, 2009: 758-759.",
        "[2] LEWIS P, PEREZ E, PIKTUS A, et al. Retrieval-augmented generation for "
        "knowledge-intensive NLP tasks[C]//Advances in Neural Information Processing "
        "Systems. 2020, 33: 9459-9474.",
        "[3] REIMERS N, GUREVYCH I. Sentence-BERT: Sentence embeddings using siamese "
        "BERT-networks[C]//Proceedings of EMNLP-IJCNLP. 2019: 3982-3992.",
        "[4] KARPUKHIN V, OGUZ B, MIN S, et al. Dense passage retrieval for open-domain "
        "question answering[C]//Proceedings of EMNLP. 2020: 6769-6781.",
        "[5] CHEN J, XIAO S, ZHANG P, et al. BGE M3-Embedding: Multi-lingual, "
        "multi-functionality, multi-granularity text embeddings through self-knowledge "
        "distillation[J]. arXiv preprint arXiv:2402.03216, 2024.",
        "[6] CHEN Z, CHEN W, SMILEY C, et al. FinQA: A dataset of numerical reasoning "
        "over financial data[C]//Proceedings of EMNLP. 2021: 3697-3711.",
        "[7] MALKOV Y A, YASHUNIN D A. Efficient and robust approximate nearest neighbor "
        "search using hierarchical navigable small world graphs[J]. IEEE Transactions "
        "on Pattern Analysis and Machine Intelligence, 2020, 42(4): 824-836.",
        "[8] KHATTAB O, ZAHARIA M. ColBERT: Efficient and effective passage search via "
        "contextualized late interaction over BERT[C]//Proceedings of SIGIR. 2020: 39-48.",
        "[9] ROBERTSON S, ZARAGOZA H. The probabilistic relevance framework: BM25 and "
        "beyond[J]. Foundations and Trends in Information Retrieval, 2009, 3(4): 333-389.",
        "[10] NOGUEIRA R, CHO K. Passage re-ranking with BERT[J]. arXiv preprint "
        "arXiv:1901.04085, 2019.",
        "[11] VASWANI A, SHAZEER N, PARMAR N, et al. Attention is all you need[C]//"
        "Advances in Neural Information Processing Systems. 2017, 30: 5998-6008.",
        "[12] DEVLIN J, CHANG M W, LEE K, et al. BERT: Pre-training of deep bidirectional "
        "transformers for language understanding[C]//Proceedings of NAACL-HLT. 2019: 4171-4186.",
        "[13] WEI J, WANG X, SCHUURMANS D, et al. Chain-of-thought prompting elicits "
        "reasoning in large language models[C]//Advances in Neural Information Processing "
        "Systems. 2022, 35: 24824-24837.",
        "[14] HU E J, SHEN Y, WALLIS P, et al. LoRA: Low-rank adaptation of large language "
        "models[J]. arXiv preprint arXiv:2106.09685, 2021.",
        "[15] LIU N F, LIN K, HEWITT J, et al. Lost in the middle: How language models use "
        "long contexts[J]. Transactions of the Association for Computational Linguistics, "
        "2024, 12: 157-173.",
    ]

    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        r = p.add_run(ref)
        set_font(r, CN_BODY, 9.5)

    doc.save(OUT)
    print(f"已生成: {OUT}")


if __name__ == "__main__":
    build()
