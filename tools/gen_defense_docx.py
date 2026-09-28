"""生成答辩准备 Word 文档"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"D:\claude_code\课程\自然语言处理")
OUT = ROOT / "答辩准备_可能的问题与参考答案.docx"

CN_FONT = "微软雅黑"
CN_BODY = "宋体"

sys.path.insert(0, str(ROOT))
from tools.gen_defense_qa import QA  # noqa: E402


def set_font(run, name=CN_BODY, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading("", level=level)
    r = h.add_run(text)
    set_font(r, CN_FONT, {1: 15, 2: 13, 3: 11.5}.get(level, 11),
             bold=True, color=RGBColor(0x1A, 0x36, 0x5D))
    return h


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.4
    r = p.add_run(text)
    set_font(r, CN_BODY, 10)
    return p


def add_shaded_para(doc, text, fill="EAF2FB", bold=False):
    """带底色的段落（用于核心结论提示框）"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.line_spacing = 1.4

    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)

    r = p.add_run(text)
    set_font(r, CN_BODY, 10, bold=bold)
    return p


def build() -> None:
    doc = Document()

    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)

    # --- 标题 ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("金融领域问答系统 · 答辩准备")
    set_font(r, CN_FONT, 18, bold=True, color=RGBColor(0x0F, 0x2A, 0x4A))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("可能的问题与参考答案 · 32 问")
    set_font(r, CN_BODY, 10.5, color=RGBColor(0x77, 0x77, 0x77))

    # --- 核心结论 ---
    add_heading(doc, "三条必须记住的核心结论", 2)
    for t in [
        "RAG 带来 80 个百分点提升（无 RAG 20% → 混合检索 100%），"
        "无 RAG 时事实型问题准确率为 0",
        "单一检索策略有结构性缺陷：向量在事实型上 Recall@1 仅 67%，"
        "BM25 在语义型上仅 80%，混合后均达 100%",
        "重排序价值显著：「营收 vs 净利润」区分度从 0.216 提升到 0.959",
    ]:
        add_shaded_para(doc, t, fill="E8F5E9")

    # --- 局限 ---
    add_heading(doc, "三个必须诚实承认的局限", 2)
    for t in [
        "语料仅 15 个 chunk（合成数据），规模偏小会降低检索难度",
        "评测用关键词匹配，属弱判定，无法评估完整性与逻辑一致性",
        "表格处理是技术债，当前作普通文本处理，向量化效果不理想",
    ]:
        add_shaded_para(doc, t, fill="FFF4E5")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("承认局限不会减分，被追问时答不上来才会减分。")
    set_font(r, CN_FONT, 10, bold=True, color=RGBColor(0xC0, 0x50, 0x00))

    doc.add_page_break()

    # --- 问题与答案 ---
    for cat_name, items in QA:
        add_heading(doc, cat_name, 1)
        for i, (question, points) in enumerate(items, 1):
            add_heading(doc, f"Q{i}. {question}", 3)
            for pt in points:
                add_bullet(doc, pt)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    doc.save(OUT)
    print(f"已生成: {OUT}")


if __name__ == "__main__":
    build()
